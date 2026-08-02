import csv
import io
import json
import uuid
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, Iterable, List

from openpyxl import load_workbook
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.config import settings
from app.models.models import Customer, Pricelist, Product


class ImportValidationError(ValueError):
    pass


def _bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None or value == "":
        return False
    return str(value).strip().lower() in {"1", "true", "yes", "sí", "si", "on"}


def parse_import_file(filename: str, content: bytes, max_rows: int = 10_000) -> List[Dict[str, Any]]:
    if len(content) > 20 * 1024 * 1024:
        raise ImportValidationError("El archivo supera el límite de 20 MB")
    lower = filename.lower()
    rows: List[Dict[str, Any]]
    if lower.endswith(".csv"):
        text = content.decode("utf-8-sig", errors="strict")
        rows = [dict(row) for row in csv.DictReader(io.StringIO(text))]
    elif lower.endswith(".xlsx"):
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        try:
            headers = [str(value).strip() if value is not None else "" for value in next(values)]
        except StopIteration as exc:
            raise ImportValidationError("El archivo XLSX está vacío") from exc
        if not all(headers) or len(set(headers)) != len(headers):
            raise ImportValidationError("La cabecera contiene nombres vacíos o duplicados")
        rows = [dict(zip(headers, row)) for row in values if any(value not in (None, "") for value in row)]
    elif lower.endswith(".json"):
        payload = json.loads(content.decode("utf-8"))
        if isinstance(payload, dict):
            payload = payload.get("items")
        if not isinstance(payload, list) or not all(isinstance(item, dict) for item in payload):
            raise ImportValidationError("El JSON debe contener una lista de objetos o una propiedad items")
        rows = payload
    else:
        raise ImportValidationError("Formato no soportado. Usa CSV, XLSX o JSON")
    if not rows:
        raise ImportValidationError("El archivo no contiene registros")
    if len(rows) > max_rows:
        raise ImportValidationError(f"El archivo contiene más de {max_rows} registros")
    return rows


class ImportService:
    SUPPORTED = {"products", "customers", "stock", "pricelists"}

    def __init__(self, db: Session):
        self.db = db

    def process(self, entity_type: str, items: Iterable[Dict[str, Any]], dry_run: bool = False) -> Dict[str, Any]:
        if entity_type not in self.SUPPORTED:
            raise ImportValidationError("Tipo de entidad no soportado")
        normalized_items = list(items)
        if not normalized_items or len(normalized_items) > 10_000:
            raise ImportValidationError("La importación debe contener entre 1 y 10.000 registros")

        errors: list[dict[str, Any]] = []
        prepared: list[Any] = []
        seen_keys: set[str] = set()

        for index, item in enumerate(normalized_items, start=1):
            try:
                if entity_type == "products":
                    sku = str(item.get("sku") or "").strip()
                    name = str(item.get("name") or "").strip()
                    if not sku or not name:
                        raise ValueError("sku y name son obligatorios")
                    if sku.lower() in seen_keys:
                        raise ValueError("SKU duplicado dentro del archivo")
                    seen_keys.add(sku.lower())
                    price = Decimal(str(item.get("base_price", "")))
                    stock = int(item.get("stock_quantity") or 0)
                    tax_rate = Decimal(str(item.get("tax_rate", settings.DEFAULT_VAT_RATE)))
                    if price < 0 or stock < 0 or tax_rate < 0 or tax_rate > 1:
                        raise ValueError("precio, stock o tipo impositivo fuera de rango")
                    if self.db.query(Product).filter(Product.sku == sku).first():
                        raise ValueError("SKU ya existente")
                    prepared.append(
                        Product(
                            id=item.get("id") or f"prod-{uuid.uuid4().hex[:12]}",
                            external_id=str(item.get("external_id") or "").strip() or None,
                            sku=sku,
                            name=name,
                            description=str(item.get("description") or "").strip() or None,
                            category=str(item.get("category") or "General").strip(),
                            attributes=item.get("attributes") if isinstance(item.get("attributes"), dict) else {},
                            source=str(item.get("source") or "import").strip(),
                            source_url=str(item.get("source_url") or "").strip() or None,
                            base_price=price,
                            tax_rate=tax_rate,
                            stock_quantity=stock,
                            reserved_stock=0,
                        )
                    )
                elif entity_type == "customers":
                    name = str(item.get("name") or "").strip()
                    email = str(item.get("email") or "").strip().lower()
                    if not name or not email or "@" not in email:
                        raise ValueError("name y un email válido son obligatorios")
                    if email in seen_keys:
                        raise ValueError("correo duplicado dentro del archivo")
                    seen_keys.add(email)
                    if self.db.query(Customer).filter(Customer.email == email).first():
                        raise ValueError("correo ya existente")
                    pricelist_id = str(item.get("pricelist_id") or "").strip() or None
                    if pricelist_id and not self.db.query(Pricelist).filter(Pricelist.id == pricelist_id).first():
                        raise ValueError("tarifa no encontrada")
                    prepared.append(
                        Customer(
                            id=item.get("id") or f"cust-{uuid.uuid4().hex[:12]}",
                            name=name,
                            email=email,
                            phone=str(item.get("phone") or "").strip() or None,
                            company=str(item.get("company") or "").strip() or None,
                            pricelist_id=pricelist_id,
                            tax_exempt=_bool(item.get("tax_exempt")),
                        )
                    )
                elif entity_type == "pricelists":
                    name = str(item.get("name") or "").strip()
                    identifier = str(item.get("id") or "").strip() or f"pl-{uuid.uuid4().hex[:12]}"
                    if not name:
                        raise ValueError("name es obligatorio")
                    if identifier.lower() in seen_keys:
                        raise ValueError("ID de tarifa duplicado dentro del archivo")
                    seen_keys.add(identifier.lower())
                    if self.db.query(Pricelist).filter(Pricelist.id == identifier).first():
                        raise ValueError("ID de tarifa ya existente")
                    discount = Decimal(str(item.get("discount_percentage", 0)))
                    if discount < 0 or discount > 100:
                        raise ValueError("descuento fuera de rango")
                    prepared.append(Pricelist(id=identifier, name=name, discount_percentage=discount))
                else:  # stock
                    sku = str(item.get("sku") or "").strip()
                    if not sku:
                        raise ValueError("sku es obligatorio")
                    if sku.lower() in seen_keys:
                        raise ValueError("SKU duplicado dentro del archivo")
                    seen_keys.add(sku.lower())
                    product = self.db.query(Product).filter(Product.sku == sku).first()
                    if not product:
                        raise ValueError("producto no encontrado")
                    quantity = int(item.get("stock_quantity"))
                    if quantity < int(product.reserved_stock):
                        raise ValueError("el stock no puede ser inferior al reservado")
                    prepared.append((product, quantity))
            except (InvalidOperation, KeyError, TypeError, ValueError, ArithmeticError) as exc:
                errors.append({"row": index, "error": str(exc)})

        if errors:
            self.db.rollback()
            return {"status": "validation_error", "valid_rows": len(prepared), "invalid_rows": len(errors), "errors": errors}
        if dry_run:
            self.db.rollback()
            return {"status": "dry_run", "valid_rows": len(prepared), "invalid_rows": 0, "errors": []}

        try:
            updated = 0
            for prepared_item in prepared:
                if entity_type == "stock":
                    product, quantity = prepared_item
                    product.stock_quantity = quantity
                    updated += 1
                else:
                    self.db.add(prepared_item)
            self.db.commit()
            return {
                "status": "success",
                "imported_count": len(prepared) if entity_type != "stock" else 0,
                "updated_count": updated,
                "errors": [],
            }
        except IntegrityError as exc:
            self.db.rollback()
            raise ImportValidationError("La importación viola una restricción de integridad o unicidad") from exc
