import csv
import hashlib
import math
import mimetypes
import re
import shutil
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional

import httpx
from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from openpyxl import load_workbook
from pypdf import PdfReader
from sklearn.decomposition import TruncatedSVD
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.preprocessing import normalize
from sqlalchemy.orm import Session

from app.config import settings
from app.models.models import RAGChunk, RAGDocument


class RAGError(ValueError):
    pass


class RAGService:
    SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".html", ".htm", ".pdf", ".docx", ".xlsx", ".png", ".jpg", ".jpeg", ".webp", ".tif", ".tiff"}

    def __init__(self, db: Session):
        self.db = db

    @staticmethod
    def _safe_roots() -> List[Path]:
        roots = []
        for root in settings.rag_allowed_roots:
            roots.append(Path(root).expanduser().resolve())
        return roots

    @classmethod
    def ensure_allowed_path(cls, file_path: str) -> Path:
        path = Path(file_path).expanduser().resolve()
        if settings.ENVIRONMENT.lower() == "testing":
            return path
        roots = cls._safe_roots()
        if not roots or not any(path == root or root in path.parents for root in roots):
            raise RAGError("La ruta del documento no pertenece al directorio autorizado")
        return path

    @staticmethod
    def _validate_archive(path: Path) -> None:
        """Evita archivos Office comprimidos con expansión desproporcionada."""
        try:
            with zipfile.ZipFile(path) as archive:
                members = archive.infolist()
                if len(members) > 10_000:
                    raise RAGError("El archivo comprimido contiene demasiados elementos")
                compressed = sum(max(member.compress_size, 1) for member in members)
                expanded = sum(member.file_size for member in members)
                if expanded > 200 * 1024 * 1024 or expanded / compressed > 100:
                    raise RAGError("El archivo comprimido supera los límites de seguridad")
        except zipfile.BadZipFile as exc:
            raise RAGError("El archivo Office está corrupto") from exc

    @staticmethod
    def _ocr_image(image: Any) -> str:
        if not settings.OCR_ENABLED:
            raise RAGError("El archivo requiere OCR y OCR_ENABLED está desactivado")
        try:
            import pytesseract
        except ImportError as exc:
            raise RAGError("El soporte OCR no está instalado") from exc
        try:
            return str(pytesseract.image_to_string(image, lang="spa+eng")).strip()
        except Exception as exc:
            raise RAGError(f"No se pudo ejecutar OCR: {exc}") from exc

    def extract_text(self, file_path: str) -> str:
        path = self.ensure_allowed_path(file_path)
        if not path.exists() or not path.is_file():
            raise RAGError("El archivo no existe")
        if path.stat().st_size <= 0:
            raise RAGError("El archivo está vacío")
        if path.stat().st_size > settings.RAG_MAX_FILE_MB * 1024 * 1024:
            raise RAGError(f"El archivo supera el límite de {settings.RAG_MAX_FILE_MB} MB")
        ext = path.suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            raise RAGError(f"Formato no soportado: {ext or 'sin extensión'}")

        if ext in {".txt", ".md"}:
            return path.read_text(encoding="utf-8", errors="replace")
        if ext == ".csv":
            rows: list[str] = []
            with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
                for row_index, row in enumerate(csv.reader(handle), start=1):
                    if row_index > 100_000:
                        raise RAGError("El CSV supera el límite de 100.000 filas")
                    rows.append(" | ".join(cell.strip() for cell in row))
            return "\n".join(rows)
        if ext in {".html", ".htm"}:
            soup = BeautifulSoup(path.read_text(encoding="utf-8", errors="replace"), "html.parser")
            for node in soup(["script", "style", "noscript", "iframe"]):
                node.decompose()
            return "\n".join(line.strip() for line in soup.get_text("\n").splitlines() if line.strip())
        if ext == ".pdf":
            try:
                reader = PdfReader(str(path))
                text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
            except Exception as exc:
                raise RAGError(f"El PDF está corrupto o no se puede leer: {exc}") from exc
            if text:
                return text
            if not settings.OCR_ENABLED:
                raise RAGError("El PDF no contiene texto extraíble; requiere OCR")
            try:
                import pypdfium2 as pdfium
                from PIL import Image

                pdf = pdfium.PdfDocument(str(path))
                pages: list[str] = []
                for index in range(len(pdf)):
                    bitmap = pdf[index].render(scale=2).to_pil()
                    if not isinstance(bitmap, Image.Image):
                        bitmap = Image.fromarray(bitmap)
                    pages.append(self._ocr_image(bitmap))
                text = "\n".join(pages).strip()
            except RAGError:
                raise
            except Exception as exc:
                raise RAGError(f"No se pudo realizar OCR del PDF: {exc}") from exc
            if not text:
                raise RAGError("OCR no encontró texto utilizable en el PDF")
            return text
        if ext == ".docx":
            self._validate_archive(path)
            try:
                document = DocxDocument(str(path))
                parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
                for table in document.tables:
                    for row in table.rows:
                        parts.append(" | ".join(cell.text.strip() for cell in row.cells))
                return "\n".join(parts)
            except RAGError:
                raise
            except Exception as exc:
                raise RAGError(f"No se pudo leer el DOCX: {exc}") from exc
        if ext == ".xlsx":
            self._validate_archive(path)
            try:
                workbook = load_workbook(str(path), read_only=True, data_only=True)
                rows: list[str] = []
                total_rows = 0
                for sheet in workbook.worksheets:
                    rows.append(f"Hoja: {sheet.title}")
                    for row in sheet.iter_rows(values_only=True):
                        total_rows += 1
                        if total_rows > 100_000:
                            raise RAGError("El XLSX supera el límite de 100.000 filas")
                        values = ["" if value is None else str(value) for value in row]
                        if any(value.strip() for value in values):
                            rows.append(" | ".join(values))
                return "\n".join(rows)
            except RAGError:
                raise
            except Exception as exc:
                raise RAGError(f"No se pudo leer el XLSX: {exc}") from exc

        try:
            from PIL import Image

            with Image.open(path) as image:
                image.verify()
            with Image.open(path) as image:
                return self._ocr_image(image.convert("RGB"))
        except RAGError:
            raise
        except Exception as exc:
            raise RAGError(f"La imagen está corrupta o no se puede procesar: {exc}") from exc

    @staticmethod
    def chunk_text(text: str, chunk_size: int = 240, overlap: int = 40) -> List[str]:
        clean = re.sub(r"\s+", " ", text).strip()
        if not clean:
            raise RAGError("El documento no contiene texto utilizable")
        words = clean.split(" ")
        chunks: list[str] = []
        step = max(1, chunk_size - overlap)
        for start in range(0, len(words), step):
            chunk = " ".join(words[start : start + chunk_size]).strip()
            if chunk:
                chunks.append(chunk)
            if start + chunk_size >= len(words):
                break
        return chunks

    @staticmethod
    def _embedding_endpoint() -> str:
        base = settings.EMBEDDING_BASE_URL.rstrip("/")
        return f"{base}/embeddings" if base.endswith("/v1") else f"{base}/v1/embeddings"

    @staticmethod
    def _validate_vectors(vectors: List[List[float]], expected_count: int) -> List[List[float]]:
        if len(vectors) != expected_count:
            raise RAGError("Número de embeddings inesperado")
        normalized: list[list[float]] = []
        for vector in vectors:
            if not isinstance(vector, list) or len(vector) != settings.EMBEDDING_DIMENSION:
                raise RAGError(
                    f"Dimensión de embedding inválida: se esperaban {settings.EMBEDDING_DIMENSION} valores"
                )
            normalized.append([float(value) for value in vector])
        return normalized

    async def _remote_embeddings(self, texts: List[str]) -> List[List[float]]:
        provider = settings.EMBEDDING_PROVIDER.lower()
        try:
            if provider == "ollama":
                vectors = []
                async with httpx.AsyncClient(timeout=60) as client:
                    for text in texts:
                        response = await client.post(
                            f"{settings.EMBEDDING_BASE_URL.rstrip('/')}/api/embeddings",
                            json={"model": settings.EMBEDDING_MODEL, "prompt": text},
                        )
                        response.raise_for_status()
                        vectors.append(response.json().get("embedding"))
                return self._validate_vectors(vectors, len(texts))
            if provider in {"openai", "openai_compatible"}:
                if not settings.EMBEDDING_API_KEY:
                    raise RAGError("Falta la credencial del proveedor de embeddings")
                async with httpx.AsyncClient(timeout=60) as client:
                    response = await client.post(
                        self._embedding_endpoint(),
                        headers={"Authorization": f"Bearer {settings.EMBEDDING_API_KEY}"},
                        json={"model": settings.EMBEDDING_MODEL, "input": texts},
                    )
                    response.raise_for_status()
                    data = response.json().get("data") or []
                return self._validate_vectors([item.get("embedding") for item in data], len(texts))
        except httpx.HTTPError as exc:
            raise RAGError(f"El proveedor de embeddings no está disponible: {exc}") from exc
        raise RAGError(f"Proveedor de embeddings no soportado: {provider}")

    @staticmethod
    def _local_tfidf_vectors(texts: List[str], query: str | None = None) -> tuple[List[List[float]], List[float] | None]:
        corpus = texts + ([query] if query is not None else [])
        vectorizer = TfidfVectorizer(lowercase=True, strip_accents="unicode", ngram_range=(1, 2), min_df=1, max_features=4096)
        matrix = vectorizer.fit_transform(corpus)
        dimension = settings.EMBEDDING_DIMENSION
        max_components = min(max(1, matrix.shape[0] - 1), max(1, matrix.shape[1] - 1), dimension)
        if max_components >= 2:
            reduced = TruncatedSVD(n_components=max_components, random_state=42).fit_transform(matrix)
        else:
            reduced = matrix.toarray()
        reduced = normalize(reduced)
        vectors = []
        for row in reduced:
            values = [float(value) for value in row.tolist()]
            vectors.append((values + [0.0] * dimension)[:dimension])
        return vectors[: len(texts)], vectors[-1] if query is not None else None

    async def add_document(self, doc_id: str, file_path: str, metadata: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
        path = self.ensure_allowed_path(file_path)
        text = self.extract_text(str(path))
        chunks = self.chunk_text(text)
        checksum = hashlib.sha256(path.read_bytes()).hexdigest()
        existing_checksum = self.db.query(RAGDocument).filter(RAGDocument.sha256 == checksum).first()
        if existing_checksum and existing_checksum.id != doc_id:
            raise RAGError(f"El documento ya existe con el ID {existing_checksum.id}")

        provider = settings.EMBEDDING_PROVIDER.lower()
        if provider == "local_tfidf":
            if settings.is_production_like:
                raise RAGError("local_tfidf solo está permitido en desarrollo y pruebas")
            vectors, _ = self._local_tfidf_vectors(chunks)
        else:
            vectors = await self._remote_embeddings(chunks)

        document_metadata = metadata or {}
        allowed_roles = document_metadata.get("allowed_roles")
        if allowed_roles is not None and (
            not isinstance(allowed_roles, list) or not all(isinstance(role, str) for role in allowed_roles)
        ):
            raise RAGError("allowed_roles debe ser una lista de roles")

        existing = self.db.query(RAGDocument).filter(RAGDocument.id == doc_id).first()
        if existing:
            self.db.query(RAGChunk).filter(RAGChunk.document_id == doc_id).delete()
            existing.file_path = str(path)
            existing.file_name = path.name
            existing.mime_type = mimetypes.guess_type(path.name)[0]
            existing.sha256 = checksum
            existing.metadata_json = document_metadata
        else:
            self.db.add(
                RAGDocument(
                    id=doc_id,
                    file_path=str(path),
                    file_name=path.name,
                    mime_type=mimetypes.guess_type(path.name)[0],
                    sha256=checksum,
                    metadata_json=document_metadata,
                )
            )
        self.db.flush()
        for index, (chunk, vector) in enumerate(zip(chunks, vectors)):
            self.db.add(
                RAGChunk(
                    id=f"{doc_id}-chunk-{index}",
                    document_id=doc_id,
                    chunk_index=index,
                    text=chunk,
                    embedding=vector,
                    metadata_json={
                        "embedding_provider": provider,
                        "embedding_model": settings.EMBEDDING_MODEL,
                        "embedding_dimension": settings.EMBEDDING_DIMENSION,
                        **document_metadata,
                    },
                )
            )
        self.db.commit()
        return {
            "doc_id": doc_id,
            "total_chunks": len(chunks),
            "file_path": str(path),
            "sha256": checksum,
            "embedding_provider": provider,
            "embedding_model": settings.EMBEDDING_MODEL,
            "embedding_dimension": settings.EMBEDDING_DIMENSION,
        }

    @staticmethod
    def _can_read(document: RAGDocument | None, role: str | None) -> bool:
        if document is None:
            return False
        allowed = (document.metadata_json or {}).get("allowed_roles")
        return not allowed or (role is not None and role in allowed)

    async def search(self, query: str, top_k: int = 3, min_score: float | None = None, role: str | None = None) -> Dict[str, Any]:
        clean_query = query.strip()
        if not clean_query:
            raise RAGError("La consulta no puede estar vacía")
        if self._looks_like_secret_request(clean_query):
            return {"abstained": True, "score": 0.0, "answer": "No puedo proporcionar secretos, credenciales ni datos privados.", "citations": []}

        chunks = self.db.query(RAGChunk).all()
        readable: list[RAGChunk] = []
        documents: dict[str, RAGDocument] = {}
        for chunk in chunks:
            document = documents.get(chunk.document_id)
            if document is None:
                document = self.db.query(RAGDocument).filter(RAGDocument.id == chunk.document_id).first()
                if document:
                    documents[chunk.document_id] = document
            if self._can_read(document, role):
                readable.append(chunk)
        if not readable:
            return {"abstained": True, "score": 0.0, "answer": "No hay documentos autorizados disponibles para responder.", "citations": []}

        provider = settings.EMBEDDING_PROVIDER.lower()
        query_vector: List[float]
        scored: list[tuple[float, RAGChunk]]
        if provider == "local_tfidf":
            document_vectors, local_query = self._local_tfidf_vectors([chunk.text for chunk in readable], clean_query)
            query_vector = local_query or []
            scored = [(self._cosine(query_vector, vector), chunk) for vector, chunk in zip(document_vectors, readable)]
        else:
            query_vector = (await self._remote_embeddings([clean_query]))[0]
            dialect = self.db.get_bind().dialect.name
            if dialect == "postgresql" and hasattr(RAGChunk.embedding, "cosine_distance"):
                distance = RAGChunk.embedding.cosine_distance(query_vector).label("distance")
                readable_ids = [chunk.id for chunk in readable]
                rows = (
                    self.db.query(RAGChunk, distance)
                    .filter(RAGChunk.id.in_(readable_ids))
                    .order_by(distance.asc())
                    .limit(max(1, min(top_k * 3, 30)))
                    .all()
                )
                scored = [(max(-1.0, min(1.0, 1.0 - float(dist))), chunk) for chunk, dist in rows]
            else:
                scored = [(self._cosine(query_vector, list(map(float, chunk.embedding))), chunk) for chunk in readable]

        scored.sort(key=lambda item: item[0], reverse=True)
        threshold = settings.RAG_MIN_SCORE if min_score is None else min_score
        selected = [(score, chunk) for score, chunk in scored[: max(1, min(top_k, 10))] if score >= threshold]
        if not selected:
            highest = scored[0][0] if scored else 0.0
            return {
                "abstained": True,
                "score": round(highest, 4),
                "answer": "No tengo evidencia suficiente en los documentos para responder con fiabilidad.",
                "citations": [],
            }
        citations = []
        evidence = []
        for score, chunk in selected:
            document = documents.get(chunk.document_id)
            citations.append(
                {
                    "doc_id": chunk.document_id,
                    "file_name": document.file_name if document else None,
                    "chunk_id": chunk.id,
                    "score": round(score, 4),
                    "snippet": chunk.text[:300],
                }
            )
            evidence.append(chunk.text)
        return {
            "abstained": False,
            "score": round(selected[0][0], 4),
            "answer": "\n\n".join(evidence[:2]),
            "citations": citations,
        }

    def delete_document(self, doc_id: str) -> bool:
        document = self.db.query(RAGDocument).filter(RAGDocument.id == doc_id).first()
        if not document:
            return False
        self.db.query(RAGChunk).filter(RAGChunk.document_id == doc_id).delete()
        file_path = Path(document.file_path)
        self.db.delete(document)
        self.db.commit()
        upload_root = Path(settings.RAG_UPLOAD_DIR).resolve()
        try:
            if file_path.exists() and (file_path == upload_root or upload_root in file_path.resolve().parents):
                file_path.unlink(missing_ok=True)
        except OSError:
            pass
        return True

    @staticmethod
    def save_upload(filename: str, content: bytes) -> Path:
        suffix = Path(filename).suffix.lower()
        if suffix not in RAGService.SUPPORTED_EXTENSIONS:
            raise RAGError("Formato de archivo no soportado")
        if not content or len(content) > settings.RAG_MAX_FILE_MB * 1024 * 1024:
            raise RAGError("Archivo vacío o superior al tamaño permitido")
        upload_dir = Path(settings.RAG_UPLOAD_DIR).resolve()
        upload_dir.mkdir(parents=True, exist_ok=True)
        digest = hashlib.sha256(content).hexdigest()
        safe_path = upload_dir / f"{digest[:24]}{suffix}"
        if not safe_path.exists():
            safe_path.write_bytes(content)
        return safe_path

    @staticmethod
    def _cosine(a: List[float], b: List[float]) -> float:
        if not a or not b or len(a) != len(b):
            return 0.0
        dot = sum(x * y for x, y in zip(a, b))
        norm_a = math.sqrt(sum(x * x for x in a))
        norm_b = math.sqrt(sum(y * y for y in b))
        return dot / (norm_a * norm_b) if norm_a and norm_b else 0.0

    @staticmethod
    def _looks_like_secret_request(query: str) -> bool:
        lowered = query.lower()
        patterns = (
            "contraseña",
            "password",
            "api key",
            "token secreto",
            "secret",
            "credencial",
            "bypass rbac",
            "superadministrador",
            "system prompt",
            "variables de entorno",
        )
        return any(pattern in lowered for pattern in patterns)
